import sqlite3
from openpyxl import load_workbook
from db import BASE, DB_PATH

DATA = BASE / "data"

def get_or_create(cur, table, name):
    if not name:
        name = "Неизвестно"
    cur.execute(f"INSERT OR IGNORE INTO {table}(name) VALUES (?)", (name,))
    row = cur.execute(f"SELECT id FROM {table} WHERE name=?", (name,)).fetchone()
    return row[0]

def get_or_create_category(cur, cat, subcat):
    if not cat:
        cat = "Без категории"
    if not subcat:
        subcat = "Без подкатегории"
    cur.execute("INSERT OR IGNORE INTO categories(name, subcategory) VALUES (?, ?)", (cat, subcat))
    row = cur.execute("SELECT id FROM categories WHERE name=? AND subcategory=?", (cat, subcat)).fetchone()
    return row[0]

def iter_rows(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = wb.active.iter_rows(values_only=True)
    next(rows, None) # пропуск заголовка
    for r in rows:
        # Пропускаем полностью пустые строки
        if r and any(cell is not None for cell in r):
            yield r

def to_float(v):
    if v is None or v == "": return 0.0
    if isinstance(v, (int, float)): return float(v)
    s = str(v).replace(",", ".")
    try: return float(s)
    except ValueError: return 0.0

# 1. Пользователи (A: Фамилия, B: Имя, C: Отчество, D: Логин, E: Роль)
def import_users(cur):
    for r in iter_rows(DATA / "Users_import.xlsx"):
        if len(r) < 5 or not r[3]: continue
        full_name = f"{r[0]} {r[1]} {r[2]}".strip()
        login = r[3]
        role = r[4]
        rid = get_or_create(cur, "roles", role)
        cur.execute(
            "INSERT OR IGNORE INTO users(login, full_name, role_id) VALUES (?,?,?)",
            (login, full_name, rid)
        )

# 2. Модели обуви (C: Изображение, D: Наименование, E: Производство, F: Описание, G: Состав, H: Цена)
def import_models(cur):
    for r in iter_rows(DATA / "Products_import.xlsx"):
        if len(r) < 8 or not r[3]: continue # r[3] это колонка D (Наименование)
        
        image = r[2]
        name  = r[3]
        manuf = r[4]
        desc  = r[5]
        comp  = r[6]
        price = r[7]
        
        # Категории из модели неизвестны (колонки A, B пустые), поэтому ставим заглушку
        cid = get_or_create_category(cur, "Обувь", "Разное") 
        mid = get_or_create(cur, "manufacturers", manuf)
        
        cur.execute(
            """INSERT OR IGNORE INTO models(name, category_id, manufacturer_id,
                                            description, composition, price, image_path)
               VALUES (?,?,?,?,?,?,?)""",
            (name, cid, mid, desc, comp, to_float(price), image)
        )

# 3. Размеры (A: Размер)
def import_sizes(cur):
    for r in iter_rows(DATA / "Sizes_import.xlsx"):
        if not r[0]: continue
        val = to_float(r[0])
        cur.execute("INSERT OR IGNORE INTO sizes(value) VALUES (?)", (val,))

# 4. Остатки (A: Наименование, B: Производство, C: Размер, D: Количество)
def import_stock(cur):
    for r in iter_rows(DATA / "Stock_Items_import.xlsx"):
        if len(r) < 4: continue
        name, size_val, qty = r[0], to_float(r[2]), int(to_float(r[3]))
        
        m = cur.execute("SELECT id FROM models WHERE name=?", (name,)).fetchone()
        s = cur.execute("SELECT id FROM sizes WHERE value=?", (size_val,)).fetchone()
        
        if not m or not s:
            continue
            
        cur.execute(
            "INSERT OR IGNORE INTO stock_items(model_id, size_id, quantity) VALUES (?,?,?)",
            (m[0], s[0], qty)
        )

# 5. Заказы (A: Номер, B: Дата, C: ФИО, D: Категория, E: Наименование, F: Производство, G: Размер, H: Кол-во, I: Цена)
def import_orders(cur):
    current_order_id = None
    
    for r in iter_rows(DATA / "Orders_import.xlsx"):
        if len(r) < 9: continue
        
        order_num = r[0]
        date      = r[1]
        full_name = r[2]
        name      = r[4]
        size_val  = to_float(r[6])
        qty       = int(to_float(r[7]))
        price     = to_float(r[8])
        
        # В Excel номер заказа и дата заполнены только в первой строке заказа
        if order_num:
            current_order_id = order_num
            
            # Ищем пользователя по ФИО (так как логина в файле заказов нет)
            u = cur.execute("SELECT id FROM users WHERE full_name=?", (full_name,)).fetchone()
            if u:
                cur.execute(
                    "INSERT OR IGNORE INTO orders(id, order_date, user_id) VALUES (?,?,?)",
                    (int(order_num), str(date), u[0])
                )
            else:
                print(f"  [!] Пользователь не найден: {full_name}")

        if current_order_id:
            # Находим ID товарной позиции
            m = cur.execute("SELECT id FROM models WHERE name=?", (name,)).fetchone()
            s = cur.execute("SELECT id FROM sizes WHERE value=?", (size_val,)).fetchone()
            
            if m and s:
                si = cur.execute("SELECT id FROM stock_items WHERE model_id=? AND size_id=?", (m[0], s[0])).fetchone()
                if si:
                    cur.execute(
                        "INSERT INTO order_items(order_id, stock_item_id, quantity, price_per_unit) VALUES (?,?,?,?)",
                        (int(current_order_id), si[0], qty, price)
                    )
                else:
                    print(f"  [!] Нет остатка для: {name} (размер {size_val})")

def main():
    c = sqlite3.connect(DB_PATH)
    cur = c.cursor()
    cur.execute("PRAGMA foreign_keys = ON")
    
    print("1. Импорт пользователей...")
    import_users(cur)
    print("2. Импорт моделей...")
    import_models(cur)
    print("3. Импорт размеров...")
    import_sizes(cur)
    print("4. Импорт остатков...")
    import_stock(cur)
    print("5. Импорт заказов...")
    import_orders(cur)
    
    c.commit()
    c.close()
    print("Импорт успешно завершен.")

if __name__ == "__main__":
    main()